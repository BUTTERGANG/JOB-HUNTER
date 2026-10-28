#!/usr/bin/env python3
"""Turn the work-history spreadsheet into a master resume the app can rank from.

Two sources are merged: the work-history spreadsheet (detail per role) and the
LinkedIn profile (dates, titles, education, certifications, and roles the sheet
never had). The result is the master resume text the app reads from
`resumes(type='master')`, keeping every bullet, skill and achievement so the tailor
has real material to reorder and rephrase — it is explicitly forbidden from
inventing anything, so anything missing here is missing from every tailored resume.

Usage:
  .venv/bin/python3 scripts/work-history-to-resume.py \
      --in  ~/Downloads/Alex_Butterfield_Work_History_Master.xlsx \
      --out-md data/master_resume.md \
      --out-csv data/work_history.csv \
      --out-xlsx ~/Downloads/Alex_Butterfield_Work_History_Master_corrected.xlsx

Output goes to `data/` by default, which is gitignored — resume content must never
be committed to this (public) repository.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

SHEET = "Work History"

# Corrections applied on top of the spreadsheet, keyed by (company, title).
# Kept here rather than hand-edited into the output so a re-import from the
# original file cannot silently lose them.
CORRECTIONS: dict[tuple[str, str], dict[str, str]] = {
    ("The Mill Top", "Event Operations & Security Staff"): {
        # Still working there — the sheet listed a single year, which reads as an
        # ended role and makes the experience look stale to a ranker or a human.
        "Year": "2022-Present",
    },
    # LinkedIn carries the public-facing title and a more precise end date.
    ("Meet Week Co.", "Event Coordinator & Media Specialist"): {
        "Position/Title": "Executive Producer",
        "Year": "2022-2025",
    },
    # LinkedIn: November 2019 – December 2020.
    ("Amazon (DIN1 Delivery Center)", "Warehouse Associate"): {"Year": "2019-2020"},
}

# Where each employer is, from the LinkedIn profile. The DOCX prints
# "title | company | location", so an empty location leaves a dangling separator.
LOCATIONS: dict[str, str] = {
    "LA Media": "Carmel, IN",
    "COMPLETE Weddings + Events": "Indianapolis, IN",
    "Moran Academy for Strength and Speed": "Carmel, IN",
    "Meet Week Co.": "Carmel, IN",
    "The Mill Top": "Noblesville, IN",
    "Quickship Brands (QSB)": "Indianapolis, IN",
    "Amazon (DIN1 Delivery Center)": "Indianapolis, IN",
    "Amazon (Return Center)": "Indianapolis, IN",
    "LA Fitness": "Fishers, IN",
    "F45 Training (Clay Terrace, Carmel)": "Carmel, IN",
    "AB Functional Fitness": "Indianapolis, IN",
    "Adkins Drapery and Blinds": "Indianapolis, IN",
    "Sherwin-Williams": "Fishers, IN",
    "Lifetime Fitness": "Fishers, IN",
}

# Facts transcribed from https://www.linkedin.com/in/alex-butterfield-7865a0204/
# Anything not verifiable in either source stays out.
LINKEDIN = {
    "profile_url": "linkedin.com/in/alex-butterfield-7865a0204",
    "location": "Carmel, Indiana",
    "about": (
        "Experience in photography/videography, logistics, audio/visual event production and "
        "coordination, Olympic weightlifting and sports performance, liquidation services, "
        "electronic replacement services, eBay refurbished resale services, and "
        "Windows/Linux/Mac/SaaS/CRM systems. Has worked in various roles and often fills the "
        "gaps in a team with the ability to adapt on the fly in high-pressure environments."
    ),
    "education": [
        {
            "school": "Fishers High School",
            "credential": "High School Diploma",
            "location": "Fishers, IN",
            "start": "2012",
            "end": "2016",
        }
    ],
    # Every one of these has an end date in the past — kept as stated, never renewed
    # silently. Whether to show an expired certification is the tailor's judgement.
    "certifications": [
        "USAW Coach Level 1 — USA Weightlifting (Dec 2020 – Dec 2021)",
        "CPR & AED — NationalCPRFoundation (Dec 2020 – Dec 2022)",
        "Google Ads Display Certification — Google (Jan 2021 – Jan 2022)",
    ],
}

# Rows to add that are not in the spreadsheet yet. Every field left as a TODO is
# deliberately not guessed — the LLC's legal name, its purpose and the role are
# facts only the owner can supply.
EXTRA_ROWS: list[dict[str, str]] = [
    {
        # On the LinkedIn profile but missing from the spreadsheet — an ongoing role,
        # so its absence made the history look both incomplete and less current.
        "Year": "2020-Present",
        "Company": "Moran Academy for Strength and Speed",
        "Position/Title": "Olympic Lifting Coach",
        "Responsibilities": (
            "Perform movement assessments to identify muscular imbalances and deficiencies in "
            "athletes' biomechanics, give verbal cues and prescribe accessory work to develop "
            "the technical execution of the Olympic lifts, coach athletes competing at local and "
            "national level events such as the American Open Series"
        ),
        "Skills Gained": (
            "Olympic weightlifting technique, movement screening, biomechanics, programming "
            "accessory work, athlete coaching, competition preparation, verbal cueing"
        ),
        "Key Achievements/Metrics": (
            "Coaches athletes through to national-level competition (American Open Series)"
        ),
        "Relevant For (Resume Type)": "Coaching, Operations",
    },
    {
        "Year": "TODO: start year-Present",
        "Company": "BUTTERFIELD LLC",
        "Position/Title": "TODO: your role/title in the LLC",
        "Responsibilities": (
            "TODO: what the LLC does and what you do in it. Not on the spreadsheet yet — "
            "fill this in and re-run, or edit it in Settings → Master Resume."
        ),
        "Skills Gained": "TODO: business operations, entity/contract management, etc.",
        "Key Achievements/Metrics": "TODO: revenue, client counts, or other metrics.",
        "Relevant For (Resume Type)": "TODO: Photography, Operations, etc.",
    },
]

HEADER = [
    "Year",
    "Company",
    "Position/Title",
    "Responsibilities",
    "Skills Gained",
    "Key Achievements/Metrics",
    "Relevant For (Resume Type)",
]


def read_rows(path: Path) -> list[dict[str, str]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"sheet {SHEET!r} not found in {path} (has: {wb.sheetnames})")
    ws = wb[SHEET]

    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    rows: list[dict[str, str]] = []
    for r in range(2, ws.max_row + 1):
        row = {}
        for i, key in enumerate(header):
            if not key:
                continue
            value = ws.cell(r, i + 1).value
            row[key] = str(value).strip() if value is not None else ""
        if row.get("Company"):
            rows.append(row)
    return rows


def apply_corrections(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    out = []
    for row in rows:
        patch = CORRECTIONS.get((row.get("Company", ""), row.get("Position/Title", "")))
        if patch:
            row = {**row, **patch}
        location = LOCATIONS.get(row.get("Company", ""))
        if location:
            row["Location"] = location
        out.append(row)
    return out


def split_phrases(value: str) -> list[str]:
    """Split a comma-separated cell into bullets, keeping short clauses intact."""
    if not value:
        return []
    parts = [p.strip(" .") for p in value.split(", ")]
    return [p for p in parts if p]


def parse_sort_key(row: dict[str, str]) -> tuple[int, int]:
    """Sort newest-first. '2023-Present' outranks '2023'; a missing year sinks."""
    year = row.get("Year", "")
    match = re.search(r"(19|20)\d{2}", year)
    start = int(match.group(0)) if match else 0
    current = 1 if re.search(r"present", year, re.I) else 0
    return (start, current)


def to_markdown(rows: list[dict[str, str]], name: str) -> str:
    lines: list[str] = []
    lines.append(f"# {name} — Master Work History")
    lines.append("")
    lines.append(f"Location: {LINKEDIN['location']}")
    lines.append(f"LinkedIn: {LINKEDIN['profile_url']}")
    lines.append(
        "Email, phone and portfolio are NOT in this file. Do not invent them — leave those "
        "contact fields empty in your output rather than making a value up."
    )
    lines.append("")
    lines.append("## In his own words (raw material for the professional summary)")
    lines.append(LINKEDIN["about"])
    lines.append("")
    lines.append(
        "Every entry below is real experience. Bullets are grouped by role in reverse "
        "chronological order; the `Resume tags` line on each role says which resume "
        "variants that role has been used for."
    )
    lines.append("")

    # Rows still carrying a TODO are placeholders, not experience. Feeding them to the
    # model would risk a tailored resume containing "TODO: your role/title", so they are
    # excluded from the resume text but kept in the CSV/XLSX so the slot stays visible.
    ordered = [r for r in sorted(rows, key=parse_sort_key, reverse=True)
               if not any("TODO" in str(v) for v in r.values())]
    for row in ordered:
        company = row.get("Company", "").strip()
        title = row.get("Position/Title", "").strip()
        year = row.get("Year", "").strip()
        lines.append(f"## {company} — {title}")
        lines.append(f"Dates: {year}")
        if row.get("Location"):
            lines.append(f"Location: {row['Location']}")
        tags = row.get("Relevant For (Resume Type)", "").strip()
        if tags:
            lines.append(f"Resume tags: {tags}")
        lines.append("")

        responsibilities = split_phrases(row.get("Responsibilities", ""))
        if responsibilities:
            lines.append("Responsibilities and duties:")
            for item in responsibilities:
                lines.append(f"- {item}")
            lines.append("")

        skills = row.get("Skills Gained", "")
        if skills:
            lines.append(f"Skills gained: {skills}")
            lines.append("")

        achievements = row.get("Key Achievements/Metrics", "")
        if achievements:
            lines.append(f"Key achievements and metrics: {achievements}")
            lines.append("")

    lines.append("## Education")
    for edu in LINKEDIN["education"]:
        parts = [edu["school"]]
        if edu.get("credential"):
            parts.append(edu["credential"])
        if edu.get("location"):
            parts.append(edu["location"])
        lines.append(f"- {' | '.join(parts)} ({edu['start']}–{edu['end']})")
    lines.append("")

    lines.append("## Certifications")
    for cert in LINKEDIN["certifications"]:
        lines.append(f"- {cert}")
    lines.append(
        "Note: each certification above has an expiry date that has already passed. They are "
        "listed exactly as issued — never extend or renew a date, and drop one entirely if it "
        "is irrelevant to the role rather than implying it is current."
    )
    lines.append("")

    # A flat skills index helps the tailor mirror a posting's terminology.
    all_skills: list[str] = []
    for row in ordered:
        all_skills.extend(split_phrases(row.get("Skills Gained", "")))
    seen: set[str] = set()
    unique_skills: list[str] = []
    for skill in all_skills:
        key = skill.lower()
        if key not in seen and len(skill) > 2:
            seen.add(key)
            unique_skills.append(skill)
    if unique_skills:
        lines.append("## Skills index (every skill evidenced above)")
        lines.append(", ".join(sorted(unique_skills, key=str.lower)))
        lines.append("")

    return "\n".join(lines)


def write_csv(rows: list[dict[str, str]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=HEADER)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in HEADER})


def write_xlsx(rows: list[dict[str, str]], path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:  # a fresh Workbook always has one, but keep the type honest
        ws = wb.create_sheet(SHEET)
    ws.title = SHEET
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(k, "") for k in HEADER])
    widths = [14, 26, 32, 60, 45, 45, 26]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--in", dest="src", required=True)
    parser.add_argument("--name", default="Alex Butterfield")
    parser.add_argument("--out-md", default="data/master_resume.md")
    parser.add_argument("--out-csv", default="data/work_history.csv")
    parser.add_argument("--out-xlsx", default=None)
    args = parser.parse_args()

    src = Path(os.path.expanduser(args.src))
    rows = apply_corrections(read_rows(src))
    for extra in EXTRA_ROWS:
        if not any(r.get("Company") == extra["Company"] for r in rows):
            # Extra rows bypass apply_corrections, so attach the location here too.
            row = dict(extra)
            if LOCATIONS.get(row.get("Company", "")):
                row["Location"] = LOCATIONS[row["Company"]]
            rows.append(row)

    md = to_markdown(rows, args.name)
    md_path = Path(args.out_md)
    md_path.parent.mkdir(parents=True, exist_ok=True)
    md_path.write_text(md, encoding="utf-8")

    write_csv(rows, Path(args.out_csv))
    if args.out_xlsx:
        write_xlsx(rows, Path(os.path.expanduser(args.out_xlsx)))

    print(f"roles: {len(rows)}")
    print(f"master resume: {md_path} ({len(md)} chars)")
    print(f"csv: {args.out_csv}")
    if args.out_xlsx:
        print(f"corrected spreadsheet: {args.out_xlsx}")
    todos = sum(1 for r in rows for v in r.values() if "TODO" in str(v))
    if todos:
        print(f"NOTE: {todos} TODO field(s) remain — they are placeholders, not content")


if __name__ == "__main__":
    main()
